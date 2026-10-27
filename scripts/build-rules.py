"""Build the in-game and editable editions of the source-derived review manual."""
from pathlib import Path
import re, html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
text=(ROOT/'docs/IMPLEMENTED-RULES.md').read_text()
blocks=[]; lines=text.splitlines(); i=0
while i<len(lines):
 line=lines[i]
 if not line.strip():i+=1;continue
 if line.startswith('#'):
  m=re.match(r'^(#+) (.*)',line);blocks.append(('heading',len(m[1]),m[2]));i+=1
 elif line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   row=[x.strip() for x in lines[i].strip('|').split('|')]
   if not all(re.match(r'^:?-+:?$',x) for x in row):rows.append(row)
   i+=1
  blocks.append(('table',rows))
 elif line.startswith('- '):blocks.append(('bullet',line[2:]));i+=1
 else:
  p=[line];i+=1
  while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','- ')):p.append(lines[i]);i+=1
  blocks.append(('paragraph',' '.join(p)))
def inline(s):return re.sub(r'\*\*(.*?)\*\*',r'<strong>\1</strong>',html.escape(s))
def anchor(s):return re.sub(r'[^a-z0-9]+','-',s.lower()).strip('-')
body=[];toc=[];opened=False;in_list=False
for b in blocks:
 if in_list and b[0]!='bullet':body.append('</ul>');in_list=False
 if b[0]=='heading':
  _,level,value=b
  if level==2:
   if opened:body.append('</section>')
   body.append('<section>');opened=True;toc.append(f'<a href="#{anchor(value)}">{html.escape(value)}</a>')
  body.append(f'<h{level} id="{anchor(value)}">{html.escape(value)}</h{level}>')
 elif b[0]=='paragraph':body.append('<p>'+inline(b[1])+'</p>')
 elif b[0]=='bullet':
  if not in_list:body.append('<ul>');in_list=True
  body.append('<li>'+inline(b[1])+'</li>')
 elif b[0]=='table':
  rows=b[1];body.append('<div class="table-scroll" tabindex="0"><table><thead><tr>'+''.join('<th>'+inline(c)+'</th>'for c in rows[0])+'</tr></thead><tbody>')
  body.extend('<tr>'+''.join('<td>'+inline(c)+'</td>'for c in row)+'</tr>'for row in rows[1:]);body.append('</tbody></table></div>')
if in_list:body.append('</ul>')
if opened:body.append('</section>')
page='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="theme-color" content="#091923"><title>K Dwarf implemented rules for review</title><style>
*{box-sizing:border-box}body{margin:0;background:#091923;color:#d8e8eb;font:17px/1.65 system-ui,sans-serif}header{position:sticky;top:0;background:#0c2230f2;backdrop-filter:blur(12px);padding:12px max(18px,calc((100vw - 1100px)/2));border-bottom:1px solid #375360;z-index:2}nav{display:flex;flex-wrap:wrap;gap:12px;align-items:center}a{color:#87dfdd}input{min-width:150px;flex:1;padding:10px;background:#071720;color:#fff;border:1px solid #66838c;border-radius:5px;font:inherit}button{padding:10px;background:#234b56;color:#fff;border:1px solid #6aa6ad;border-radius:5px;cursor:pointer}main{max-width:1100px;margin:auto;padding:28px 22px 80px}h1{font-size:clamp(1.8rem,5vw,3rem);line-height:1.16;color:#fff}h2{color:#9fe8e3;margin:48px 0 16px;line-height:1.25;scroll-margin-top:140px}h3{margin-top:28px;color:#c5e9ed;scroll-margin-top:140px}p,li{max-width:88ch}li{margin:12px 0}.table-scroll{overflow-x:auto;margin:20px 0;max-width:100%;border:1px solid #365562;border-radius:6px}table{border-collapse:collapse;width:100%;font-size:.9rem}th,td{padding:12px;text-align:left;border:1px solid #365562;vertical-align:top}th{background:#1d3f4d;color:#fff}tr:nth-child(even){background:#102630}details{padding:16px;border:1px solid #365562;border-radius:6px}details a{display:block;margin:8px 0}summary{cursor:pointer}.result{font-size:.9rem;margin:4px 0 0;color:#a3c2ca}[hidden]{display:none!important}@media print{body{background:white;color:black;font-size:10pt}header,details{display:none}main{max-width:none;padding:0}h1,h2,h3,a{color:black}section{display:block!important}th{background:#e9eef0;color:black}tr:nth-child(even){background:#f5f6f7}.table-scroll{overflow:visible}table{font-size:8.5pt}h2,h3{break-after:avoid}tr{break-inside:avoid}}
</style><header><nav><a href="./">Back to game</a><input id="search" type="search" placeholder="Find a rule or ability" aria-label="Search implemented rules"><button id="print">Print</button><a href="docs/K-Dwarf-Implemented-Rules.docx" download>Word</a><a href="docs/IMPLEMENTED-RULES.md" download>Markdown</a></nav><p class="result" id="results" aria-live="polite">Version 0.3.0 · reconstructed from the executable game only</p></header><main><details><summary>Contents</summary>'''+''.join(toc)+'''</details>'''+''.join(body)+'''</main><script>const input=document.getElementById('search'),sections=[...document.querySelectorAll('main>section')];input.addEventListener('input',()=>{const q=input.value.trim().toLowerCase();for(const s of sections)s.hidden=!!q&&!s.textContent.toLowerCase().includes(q);document.getElementById('results').textContent=q?sections.filter(s=>!s.hidden).length+' matching sections':'Version 0.3.0 · reconstructed from the executable game only';});document.getElementById('print').onclick=()=>{input.value='';input.dispatchEvent(new Event('input'));window.print();};</script></html>'''
(ROOT/'rules.html').write_text(page)
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(27)
for name,size in [('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
 d.styles[name].font.size=Pt(size);d.styles[name].paragraph_format.space_before=Pt(16);d.styles[name].paragraph_format.space_after=Pt(7)
header=sec.header.paragraphs[0];header.text='K DWARF   /   IMPLEMENTED RULES   /   REVIEW EDITION';header.style='Caption';header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Version 0.3.0   •   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def rich(p,s):
 for j,t in enumerate(re.split(r'\*\*(.*?)\*\*',s)):
  r=p.add_run(t);r.bold=bool(j%2)
for b in blocks:
 if b[0]=='heading':
  _,level,value=b;d.add_paragraph(value,'Title'if level==1 else 'Heading '+str(level-1))
 elif b[0]=='paragraph':rich(d.add_paragraph(),b[1])
 elif b[0]=='bullet':rich(d.add_paragraph(style='List Bullet'),b[1])
 else:
  rows=b[1];n=len(rows[0]);t=d.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  if rows[0][0]=='Ability':widths=[1.55,.7,4.85]
  elif rows[0][0]=='Weapon':widths=[3.5,.7,.7,1.05,1.15]
  elif rows[0][0]=='Armor':widths=[2.0,.55,.55,1.55,.75,.75]
  elif rows[0][0]=='Feature':widths=[2.2,.8,1.25,2.85]
  elif rows[0][0]=='Program':widths=[1.2,5.9]
  elif rows[0][0]=='Smoke successes':widths=[.85,4.35,1.9]
  else:widths=[7.1/n]*n
  factor=7.1/sum(widths);widths=[w*factor for w in widths]
  for col,w in zip(t.columns,widths):col.width=Inches(w)
  for k,row in enumerate(rows):
   tr=t.add_row();trPr=tr._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
   if k==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
   for j,(c,v)in enumerate(zip(tr.cells,row)):
    c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;p=c.paragraphs[0];p.paragraph_format.space_after=p.paragraph_format.space_before=Pt(1.5);p.paragraph_format.line_spacing=1.05;rich(p,v)
    for r in p.runs:r.font.size=Pt(9.5);r.bold=k==0
    if len(v)<12:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DFEAF0'if k==0 else 'F7F9FA'if k%2==0 else 'FFFFFF');pr.append(sh)
    mar=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  t._tbl.tblPr.append(borders);d.add_paragraph().paragraph_format.space_after=Pt(1)
d.core_properties.title='K Dwarf Implemented Game Rules';d.core_properties.subject='Executable game rules for Nathan Culver review';d.core_properties.version='0.3.0'
# Remove inherited title rules from the base Word styles.
for element in d.element.xpath('//w:pBdr') + d.styles.element.xpath('//w:pBdr'):
 element.getparent().remove(element)
d.save(ROOT/'docs/K-Dwarf-Implemented-Rules.docx')
print('Built rules.html and editable Word manual')
